#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl


def _run_once(input_path: Path, workdir: Path) -> subprocess.CompletedProcess:
    worker = Path(__file__).with_name("determinism_worker.py").resolve()
    return subprocess.run(
        [sys.executable, str(worker), "--input", str(input_path), "--output", str(workdir)],
        capture_output=True,
        text=True,
        check=False,
    )


def _hash_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _xlsx_logical_signature(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    sig = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sig[sheet_name] = [
            tuple("" if c is None else str(c) for c in row)
            for row in ws.iter_rows(values_only=True)
        ]
    return sig


def _compare(run_a: Path, run_b: Path) -> dict:
    a_files = sorted(p.relative_to(run_a) for p in run_a.rglob("*") if p.is_file())
    b_files = sorted(p.relative_to(run_b) for p in run_b.rglob("*") if p.is_file())

    result = {
        "files_only_in_a": [str(p) for p in sorted(set(a_files) - set(b_files))],
        "files_only_in_b": [str(p) for p in sorted(set(b_files) - set(a_files))],
        "file_results": {},
        "passed": True,
    }

    if result["files_only_in_a"] or result["files_only_in_b"]:
        result["passed"] = False

    for rel in sorted(set(a_files) & set(b_files)):
        pa, pb = run_a / rel, run_b / rel
        if rel.suffix.lower() == ".xlsx":
            sa, sb = _xlsx_logical_signature(pa), _xlsx_logical_signature(pb)
            ok = sa == sb
            result["file_results"][str(rel)] = {
                "mode": "xlsx_logical",
                "sheets_a": list(sa.keys()),
                "sheets_b": list(sb.keys()),
                "ok": ok,
            }
        else:
            ha, hb = _hash_file(pa), _hash_file(pb)
            ok = ha == hb
            result["file_results"][str(rel)] = {
                "mode": "hash",
                "hash_a": ha,
                "hash_b": hb,
                "ok": ok,
            }
        if not ok:
            result["passed"] = False

    return result


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args()

    input_path = Path(args.input).resolve()
    if not input_path.exists():
        print(f"input not found: {input_path}", file=sys.stderr)
        return 2

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        run_a, run_b = root / "run_a", root / "run_b"
        run_a.mkdir()
        run_b.mkdir()

        proc_a = _run_once(input_path, run_a)
        proc_b = _run_once(input_path, run_b)

        if proc_a.returncode != 0 or proc_b.returncode != 0:
            print("pipeline run failed", file=sys.stderr)
            print(f"run_a rc={proc_a.returncode} stderr={proc_a.stderr[:1000]}", file=sys.stderr)
            print(f"run_b rc={proc_b.returncode} stderr={proc_b.stderr[:1000]}", file=sys.stderr)
            return 3

        cmp = _compare(run_a, run_b)
        out = {
            "input": str(input_path),
            "worker": "tools/determinism_worker.py",
            "passed": cmp["passed"],
            "files_only_in_a": cmp["files_only_in_a"],
            "files_only_in_b": cmp["files_only_in_b"],
            "file_results": cmp["file_results"],
        }

        if args.json:
            print(json.dumps(out, indent=2, sort_keys=True))
        else:
            print(f"determinism: {'PASS' if out['passed'] else 'FAIL'}")
            for f, r in out["file_results"].items():
                print(f"  {f}: {'ok' if r['ok'] else 'DIFF'} ({r['mode']})")

        return 0 if out["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
